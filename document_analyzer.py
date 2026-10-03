"""Google ADK document analyzer for extracting user details from Excel workbooks."""

from __future__ import annotations

import argparse
import asyncio
import json
import os
from pathlib import Path
from typing import Any

import yaml
from google.adk.agents import Agent
from google.adk.runners import Runner
from google.adk.sessions import InMemorySessionService
from google.genai import types
from openpyxl import load_workbook

CONFIG_PATH = Path(__file__).with_name("document_analyzer.yaml")
MAX_CELL_CHARS = 20_000


def load_excel_document(file_path: str) -> dict[str, Any]:
    """Load workbook values and sheet metadata for the ADK agent."""
    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Excel file not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Expected an .xlsx or .xlsm Excel file.")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for row in worksheet.iter_rows(values_only=True):
                values = [None if value is None else str(value)[:MAX_CELL_CHARS] for value in row]
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            sheets.append({"name": worksheet.title, "rows": rows})
    finally:
        workbook.close()

    return {"source": str(path), "sheets": sheets}


def load_config() -> dict[str, Any]:
    """Read the agent behavior and extraction schema from YAML."""
    with CONFIG_PATH.open(encoding="utf-8") as config_file:
        return yaml.safe_load(config_file)


def build_agent() -> Agent:
    """Build the ADK document analyzer from the YAML definition."""
    config = load_config()
    return Agent(
        name=config["name"],
        model=os.getenv("GOOGLE_ADK_MODEL", config["model"]),
        description=config["description"],
        instruction=config["instruction"],
        tools=[load_excel_document],
    )


async def analyze_excel(file_path: str, question: str | None = None) -> str:
    """Run the ADK analyzer against an Excel workbook and return its final text."""
    document = load_excel_document(file_path)
    config = load_config()
    agent = build_agent()
    runner = Runner(
        app_name=config["app_name"],
        agent=agent,
        session_service=InMemorySessionService(),
    )
    user_request = question or config["default_request"]
    prompt = (
        f"Uploaded Excel document data:\n{json.dumps(document, ensure_ascii=False, default=str)}\n\n"
        f"User request: {user_request}"
    )
    session_id = "excel-analysis"
    user_id = "document-user"
    await runner.session_service.create_session(
        app_name=config["app_name"],
        user_id=user_id,
        session_id=session_id,
    )

    final_text = ""
    message = types.Content(role="user", parts=[types.Part(text=prompt)])
    async for event in runner.run_async(
        user_id=user_id,
        session_id=session_id,
        new_message=message,
    ):
        if event.is_final_response() and event.content and event.content.parts:
            final_text = event.content.parts[0].text or final_text
    return final_text


def response_to_json(response: str) -> Any:
    """Convert the model response, including fenced JSON, into JSON data."""
    cleaned = response.strip()
    if cleaned.startswith("```"):
        lines = cleaned.splitlines()
        if lines and lines[0].strip().startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        cleaned = "\n".join(lines).strip()
    return json.loads(cleaned)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract user details from an Excel file with Google ADK.")
    parser.add_argument("excel_file", type=Path, help="Path to the uploaded .xlsx or .xlsm file")
    parser.add_argument("--question", help="Optional question or extraction request")
    parser.add_argument("--output", type=Path, help="JSON output path; defaults to the Excel filename with .json")
    args = parser.parse_args()
    response = asyncio.run(analyze_excel(str(args.excel_file), args.question))
    output_path = args.output or args.excel_file.with_suffix(".json")
    output_path.write_text(
        json.dumps(response_to_json(response), indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(f"JSON output written to: {output_path.resolve()}")


if __name__ == "__main__":
    main()
